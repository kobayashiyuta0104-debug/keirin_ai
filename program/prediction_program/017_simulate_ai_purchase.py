# -*- coding: utf-8 -*-
"""
競輪AI：配当クラス予想の購入シミュレーション

【目的】
AIが予想した「配当クラス」に入っている3連単を全て100円ずつ購入した場合の
購入額・払戻・収支・回収率を調査する。

調査内容
① 全レースでAI予想クラス内の買い目を全購入した場合
② AI確信度別の収支
③ AI予想クラス別の収支
④ AI確信度 × AI予想クラスの収支
⑤ 確信度の「以上」条件別の収支
⑥ レースごとの購入明細

【入力】
prediction CSV:
    training_prediction(2026.1.1~2026.8.18).csv

odds CSV:
    historical_odds_2026.1.1~2026.8.18.csv

【重要】
オッズCSVの値は「100円あたりの倍率」として扱う。
例：255.4 → 100円購入なら払戻25,540円。

1レースにつき、
「AI予想クラスの範囲に入る3連単」を全て100円購入する。
実際の1-2-3着が購入対象なら、その買い目の払戻を受け取る。
"""

from pathlib import Path
from itertools import permutations
import re
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter


# ============================================================
# ① 設定
# ============================================================

PREDICTION_FILE = Path(
    r"C:\競輪AI\csv\training"
    r"\training_prediction(2026.1.1~2026.8.18).csv"
)

ODDS_FILE = Path(
    r"C:\競輪AI\csv\historical_date\historical_odds"
    r"\historical_odds_2026.1.1~2026.8.18.csv"
)

OUTPUT_DIR = Path(
    r"C:\競輪AI\csv\analysis\017_simulate_ai_purchase"
)

BET_YEN = 100

# 確信度をこの区分で集計
CONFIDENCE_BINS = [
    (0, 9, "0～9%"),
    (10, 14, "10～14%"),
    (15, 19, "15～19%"),
    (20, 24, "20～24%"),
    (25, 29, "25～29%"),
    (30, 34, "30～34%"),
    (35, 39, "35～39%"),
    (40, 100, "40%以上"),
]

# 「確信度が○％以上」の追加分析
CONFIDENCE_THRESHOLDS = [0, 15, 20, 25, 30, 35, 40]


# ============================================================
# ② AI予想クラス → オッズ範囲
# ============================================================

CLASS_RANGE = {
    "0～1,999円": (0, 1999.999999),
    "2,000～4,999円": (2000, 4999.999999),
    "5,000～9,999円": (5000, 9999.999999),
    "10,000～19,999円": (10000, 19999.999999),
    "20,000～29,999円": (20000, 29999.999999),
    "30,000～49,999円": (30000, 49999.999999),
    "50,000～99,999円": (50000, 99999.999999),
    "100,000円以上": (100000, float("inf")),
}

CLASS_ORDER = [
    "0～1,999円",
    "2,000～4,999円",
    "5,000～9,999円",
    "10,000～19,999円",
    "20,000～29,999円",
    "30,000～49,999円",
    "50,000～99,999円",
    "100,000円以上",
]


# ============================================================
# ③ 504通りの3連単列
# ============================================================

TRIFECTA_COLUMNS = [
    f"{a}-{b}-{c}"
    for a, b, c in permutations(range(1, 10), 3)
]


# ============================================================
# ④ 共通関数
# ============================================================

def read_csv_auto(path):
    """UTF-8系CSVを読み込む"""
    try:
        return pd.read_csv(path, encoding="utf-8-sig", low_memory=False)
    except UnicodeDecodeError:
        return pd.read_csv(path, encoding="cp932", low_memory=False)


def parse_percent(value):
    """'22%' → 22.0"""
    if pd.isna(value):
        return None

    s = str(value).strip().replace("%", "").replace("％", "")

    if not s:
        return None

    try:
        return float(s)
    except ValueError:
        return None


def normalize_race_no(value):
    """1R / 01R / 1 → 1R"""
    if pd.isna(value):
        return ""

    s = str(value).strip().replace("Ｒ", "R")

    m = re.search(r"(\d+)", s)
    if not m:
        return s

    return f"{int(m.group(1))}R"


def normalize_car(value):
    """車番を整数にする"""
    if pd.isna(value):
        return None

    try:
        return int(float(value))
    except (ValueError, TypeError):
        return None


def get_actual_combo(row):
    """予想CSVの1着・2着・3着から実際の3連単を作る"""
    a = normalize_car(row.get("１着"))
    b = normalize_car(row.get("２着"))
    c = normalize_car(row.get("３着"))

    if a is None or b is None or c is None:
        return None

    if len({a, b, c}) != 3:
        return None

    if not (1 <= a <= 9 and 1 <= b <= 9 and 1 <= c <= 9):
        return None

    return f"{a}-{b}-{c}"


def get_class_range(ai_class):
    return CLASS_RANGE.get(str(ai_class).strip())


def calc_class_summary(df):
    """クラス別集計"""
    rows = []

    for cls in CLASS_ORDER:
        x = df[df["AI予想"] == cls]

        if len(x) == 0:
            rows.append({
                "AI予想クラス": cls,
                "レース数": 0,
                "購入レース数": 0,
                "購入点数": 0,
                "購入金額": 0,
                "的中数": 0,
                "払戻合計": 0,
                "収支": 0,
                "回収率": 0,
                "的中率": 0,
                "1レース平均購入額": 0,
                "1レース平均購入点数": 0,
            })
            continue

        purchase_races = int((x["購入点数"] > 0).sum())
        bet_count = int(x["購入点数"].sum())
        bet_money = int(x["購入金額"].sum())
        hit_count = int((x["的中判定_sim"] == "○").sum())
        payout = int(x["払戻シミュレーション"].sum())
        profit = payout - bet_money

        rows.append({
            "AI予想クラス": cls,
            "レース数": len(x),
            "購入レース数": purchase_races,
            "購入点数": bet_count,
            "購入金額": bet_money,
            "的中数": hit_count,
            "払戻合計": payout,
            "収支": profit,
            "回収率": round(payout / bet_money * 100, 2) if bet_money else 0,
            "的中率": round(hit_count / purchase_races * 100, 2) if purchase_races else 0,
            "1レース平均購入額": round(bet_money / len(x), 2),
            "1レース平均購入点数": round(bet_count / len(x), 2),
        })

    return pd.DataFrame(rows)


def calc_confidence_summary(df):
    """確信度区分別集計"""
    rows = []

    for low, high, label in CONFIDENCE_BINS:
        x = df[
            (df["AI確信度数値"] >= low) &
            (df["AI確信度数値"] <= high)
        ]

        purchase_races = int((x["購入点数"] > 0).sum())
        bet_count = int(x["購入点数"].sum())
        bet_money = int(x["購入金額"].sum())
        hit_count = int((x["的中判定_sim"] == "○").sum())
        payout = int(x["払戻シミュレーション"].sum())
        profit = payout - bet_money

        rows.append({
            "AI確信度": label,
            "レース数": len(x),
            "購入レース数": purchase_races,
            "購入点数": bet_count,
            "購入金額": bet_money,
            "的中数": hit_count,
            "払戻合計": payout,
            "収支": profit,
            "回収率": round(payout / bet_money * 100, 2) if bet_money else 0,
            "的中率": round(hit_count / purchase_races * 100, 2) if purchase_races else 0,
            "1レース平均購入額": round(bet_money / len(x), 2) if len(x) else 0,
            "1レース平均購入点数": round(bet_count / len(x), 2) if len(x) else 0,
        })

    return pd.DataFrame(rows)


def calc_threshold_summary(df):
    """確信度○％以上だけ購入した場合"""
    rows = []

    for threshold in CONFIDENCE_THRESHOLDS:
        x = df[df["AI確信度数値"] >= threshold]

        purchase_races = int((x["購入点数"] > 0).sum())
        bet_count = int(x["購入点数"].sum())
        bet_money = int(x["購入金額"].sum())
        hit_count = int((x["的中判定_sim"] == "○").sum())
        payout = int(x["払戻シミュレーション"].sum())
        profit = payout - bet_money

        rows.append({
            "確信度条件": f"{threshold}%以上",
            "レース数": len(x),
            "購入レース数": purchase_races,
            "購入点数": bet_count,
            "購入金額": bet_money,
            "的中数": hit_count,
            "払戻合計": payout,
            "収支": profit,
            "回収率": round(payout / bet_money * 100, 2) if bet_money else 0,
            "的中率": round(hit_count / purchase_races * 100, 2) if purchase_races else 0,
            "1レース平均購入額": round(bet_money / len(x), 2) if len(x) else 0,
            "1レース平均購入点数": round(bet_count / len(x), 2) if len(x) else 0,
        })

    return pd.DataFrame(rows)


def calc_confidence_class_summary(df):
    """確信度 × AI予想クラス"""
    rows = []

    for low, high, conf_label in CONFIDENCE_BINS:
        for cls in CLASS_ORDER:
            x = df[
                (df["AI確信度数値"] >= low) &
                (df["AI確信度数値"] <= high) &
                (df["AI予想"] == cls)
            ]

            if len(x) == 0:
                continue

            purchase_races = int((x["購入点数"] > 0).sum())
            bet_count = int(x["購入点数"].sum())
            bet_money = int(x["購入金額"].sum())
            hit_count = int((x["的中判定_sim"] == "○").sum())
            payout = int(x["払戻シミュレーション"].sum())
            profit = payout - bet_money

            rows.append({
                "AI確信度": conf_label,
                "AI予想クラス": cls,
                "レース数": len(x),
                "購入レース数": purchase_races,
                "購入点数": bet_count,
                "購入金額": bet_money,
                "的中数": hit_count,
                "払戻合計": payout,
                "収支": profit,
                "回収率": round(payout / bet_money * 100, 2) if bet_money else 0,
                "的中率": round(hit_count / purchase_races * 100, 2) if purchase_races else 0,
                "1レース平均購入額": round(bet_money / len(x), 2),
                "1レース平均購入点数": round(bet_count / len(x), 2),
            })

    return pd.DataFrame(rows)


# ============================================================
# ⑤ メイン処理
# ============================================================

def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    print("=" * 70)
    print("競輪AI 購入シミュレーション開始")
    print("=" * 70)

    if not PREDICTION_FILE.exists():
        print(f"[ERROR] 予想CSVがありません:")
        print(PREDICTION_FILE)
        return

    if not ODDS_FILE.exists():
        print(f"[ERROR] オッズCSVがありません:")
        print(ODDS_FILE)
        return

    print("\n[1] CSV読み込み")
    pred = read_csv_auto(PREDICTION_FILE)
    odds = read_csv_auto(ODDS_FILE)

    print(f"予想CSV: {len(pred):,}レース")
    print(f"オッズCSV: {len(odds):,}レース")

    required_pred = [
        "レースキー",
        "AI予想",
        "AI確信度",
        "１着",
        "２着",
        "３着",
    ]

    required_odds = ["race_key"] + TRIFECTA_COLUMNS

    missing_pred = [c for c in required_pred if c not in pred.columns]
    missing_odds = [c for c in required_odds if c not in odds.columns]

    if missing_pred:
        print("\n[ERROR] 予想CSVに必要な列がありません:")
        print(missing_pred)
        return

    if missing_odds:
        print("\n[ERROR] オッズCSVに必要な列がありません:")
        print(missing_odds[:20])
        return

    # race_keyで結合
    odds_small = odds[required_odds].copy()
    odds_small = odds_small.drop_duplicates(subset=["race_key"])

    df = pred.merge(
        odds_small,
        left_on="レースキー",
        right_on="race_key",
        how="left",
        indicator=True,
    )

    print(f"\n[2] race_key結合")
    print(f"結合後: {len(df):,}レース")
    print(f"オッズあり: {(df['_merge'] == 'both').sum():,}")
    print(f"オッズなし: {(df['_merge'] != 'both').sum():,}")

    # 確信度
    df["AI確信度数値"] = df["AI確信度"].apply(parse_percent)

    # 実際の3連単
    df["実際買い目"] = df.apply(get_actual_combo, axis=1)

    # 初期値
    df["購入点数"] = 0
    df["購入金額"] = 0
    df["払戻シミュレーション"] = 0
    df["的中判定_sim"] = ""
    df["購入買い目"] = ""

    # 1レースずつ処理
    print("\n[3] 購入シミュレーション")

    processed = 0
    skipped_no_class = 0
    skipped_no_result = 0
    total_bets = 0
    total_money = 0
    total_payout = 0

    for idx, row in df.iterrows():
        ai_class = row["AI予想"]
        class_range = get_class_range(ai_class)

        if class_range is None:
            skipped_no_class += 1
            continue

        if pd.isna(row["AI確信度数値"]):
            continue

        actual_combo = row["実際買い目"]

        # オッズが存在する買い目だけ対象
        selected = []

        low, high = class_range

        for combo in TRIFECTA_COLUMNS:
            value = row[combo]

            if pd.isna(value):
                continue

            try:
                odds_value = float(value)
            except (ValueError, TypeError):
                continue

            if odds_value <= 0:
                continue

            # オッズ × 100円 = 払戻額
            payout_yen = odds_value * BET_YEN

            if low <= payout_yen <= high:
                selected.append((combo, odds_value, payout_yen))

        if not actual_combo:
            skipped_no_result += 1
            # 購入額自体は存在するが、結果不明なので
            # 集計から除外できるよう購入情報だけ保存
            df.at[idx, "購入点数"] = len(selected)
            df.at[idx, "購入金額"] = len(selected) * BET_YEN
            df.at[idx, "購入買い目"] = ",".join(x[0] for x in selected)
            continue

        bet_count = len(selected)
        bet_money = bet_count * BET_YEN

        payout = 0

        for combo, odds_value, payout_yen in selected:
            if combo == actual_combo:
                payout = int(round(payout_yen))
                break

        hit = "○" if payout > 0 else "×"

        df.at[idx, "購入点数"] = bet_count
        df.at[idx, "購入金額"] = bet_money
        df.at[idx, "払戻シミュレーション"] = payout
        df.at[idx, "的中判定_sim"] = hit
        df.at[idx, "購入買い目"] = ",".join(x[0] for x in selected)

        processed += 1
        total_bets += bet_count
        total_money += bet_money
        total_payout += payout

    # 結果が存在しないレースを集計から除外するためのフラグ
    df["集計対象"] = df["実際買い目"].notna() & df["AI確信度数値"].notna()

    analysis_df = df[df["集計対象"]].copy()

    # ========================================================
    # ⑥ 全体集計
    # ========================================================

    race_count = len(analysis_df)
    purchase_races = int((analysis_df["購入点数"] > 0).sum())
    bet_count = int(analysis_df["購入点数"].sum())
    bet_money = int(analysis_df["購入金額"].sum())
    hit_count = int((analysis_df["的中判定_sim"] == "○").sum())
    payout = int(analysis_df["払戻シミュレーション"].sum())
    profit = payout - bet_money

    overall = pd.DataFrame([{
        "対象レース数": race_count,
        "購入レース数": purchase_races,
        "購入点数": bet_count,
        "購入金額": bet_money,
        "的中数": hit_count,
        "払戻合計": payout,
        "収支": profit,
        "回収率": round(payout / bet_money * 100, 2) if bet_money else 0,
        "的中率": round(hit_count / purchase_races * 100, 2) if purchase_races else 0,
        "1レース平均購入額": round(bet_money / race_count, 2) if race_count else 0,
        "1レース平均購入点数": round(bet_count / race_count, 2) if race_count else 0,
    }])

    # ========================================================
    # ⑦ 各種集計
    # ========================================================

    print("\n[4] 各種集計")

    confidence_summary = calc_confidence_summary(analysis_df)
    class_summary = calc_class_summary(analysis_df)
    confidence_class_summary = calc_confidence_class_summary(analysis_df)
    threshold_summary = calc_threshold_summary(analysis_df)

    # ========================================================
    # ⑧ Excel保存
    # ========================================================

    # レース単位の明細
    detail_columns = [
        "レースキー",
        "日付",
        "競輪場",
        "レース\\n番号",
        "AI予想",
        "AI確信度",
        "AI確信度数値",
        "実際買い目",
        "三連単\\n払戻",
        "購入点数",
        "購入金額",
        "払戻シミュレーション",
        "的中判定_sim",
        "購入買い目",
    ]

    detail_columns = [c for c in detail_columns if c in df.columns]
    details = df[detail_columns].copy()

    # データ品質確認
    quality = pd.DataFrame([{
        "予想CSVレース数": len(pred),
        "オッズCSVレース数": len(odds),
        "race_key結合後": len(df),
        "オッズあり": int((df["_merge"] == "both").sum()),
        "オッズなし": int((df["_merge"] != "both").sum()),
        "AI予想クラス不明": skipped_no_class,
        "実際結果なし": skipped_no_result,
        "最終分析対象": len(analysis_df),
    }])

    # Excelファイル1つに全結果をまとめる
    output_excel = OUTPUT_DIR / "017_simulate_ai_purchase_result.xlsx"

    with pd.ExcelWriter(output_excel, engine="openpyxl") as writer:
        overall.to_excel(writer, sheet_name="全体集計", index=False)
        confidence_summary.to_excel(writer, sheet_name="確信度別", index=False)
        class_summary.to_excel(writer, sheet_name="クラス別", index=False)
        confidence_class_summary.to_excel(
            writer, sheet_name="確信度×クラス", index=False
        )
        threshold_summary.to_excel(
            writer, sheet_name="確信度以上", index=False
        )
        details.to_excel(writer, sheet_name="レース明細", index=False)
        quality.to_excel(writer, sheet_name="データ品質", index=False)

    # Excelの見やすさを調整
    wb = load_workbook(output_excel)

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        # ヘッダー
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # 列幅
        for col_cells in ws.columns:
            max_length = 0
            col_letter = get_column_letter(col_cells[0].column)

            for cell in col_cells:
                if cell.value is not None:
                    length = len(str(cell.value))
                    if length > max_length:
                        max_length = length

            ws.column_dimensions[col_letter].width = min(max(max_length + 2, 10), 35)

        # 金額・率を見やすくする
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, (int, float)):
                    header = ws.cell(1, cell.column).value

                    if header in [
                        "購入金額",
                        "払戻合計",
                        "収支",
                        "払戻シミュレーション",
                        "1レース平均購入額",
                    ]:
                        cell.number_format = '#,##0'

                    elif header in [
                        "回収率",
                        "的中率",
                    ]:
                        cell.number_format = '0.00'

    wb.save(output_excel)

    print(f"\\nExcel保存完了: {output_excel}")

    # ========================================================
    # ⑨ 画面表示
    # ========================================================

    print("\n" + "=" * 70)
    print("シミュレーション結果")
    print("=" * 70)

    print(f"対象レース数     : {race_count:,}")
    print(f"購入レース数     : {purchase_races:,}")
    print(f"購入点数         : {bet_count:,}")
    print(f"購入金額         : {bet_money:,} 円")
    print(f"的中数           : {hit_count:,}")
    print(f"払戻合計         : {payout:,} 円")
    print(f"収支             : {profit:+,} 円")
    print(f"回収率           : {payout / bet_money * 100:.2f}%" if bet_money else "回収率           : 0%")
    print(f"的中率           : {hit_count / purchase_races * 100:.2f}%" if purchase_races else "的中率           : 0%")
    print(f"1レース平均購入 : {bet_money / race_count:.2f} 円" if race_count else "1レース平均購入 : 0 円")
    print(f"1レース平均点数 : {bet_count / race_count:.2f} 点" if race_count else "1レース平均点数 : 0 点")

    print("\n--- 確信度別 ---")
    print(confidence_summary.to_string(index=False))

    print("\n--- AI予想クラス別 ---")
    print(class_summary.to_string(index=False))

    print("\n--- 確信度以上条件 ---")
    print(threshold_summary.to_string(index=False))

    print("\n--- 確信度 × AI予想クラス ---")
    print(confidence_class_summary.to_string(index=False))

    print("\n" + "=" * 70)
    print("保存完了")
    print(OUTPUT_DIR)
    print("=" * 70)


if __name__ == "__main__":
    main()
